"""
PDF to HWPX Full Conversion Pipeline Engine
Integrates:
  1. pdf2docx (Converter): PDF -> DOCX layout reconstruction (tables, tabs, cells)
  2. docx_to_hwpx: DOCX -> OWPML HWPX package generation
  3. Metadata & Document Statistics extraction
"""

import os
import tempfile
import fitz  # PyMuPDF for fast page counting & fallback text inspection
from typing import Dict, Any, Tuple
from pdf2docx import Converter
from backend.docx_to_hwpx import convert_docx_to_hwpx

class PdfToHwpxPipeline:
    """Full pipeline for high-fidelity PDF to HWPX conversion."""

    def __init__(self, template_path: str = None):
        self.template_path = template_path

    def convert_pdf_to_hwpx(self, pdf_file_path: str,
                            hwpx_out_path: str = None) -> Tuple[bytes, Dict[str, Any]]:
        """
        Executes the 2-step pipeline:
          Step 1: pdf2docx converts PDF to DOCX
          Step 2: docx_to_hwpx converts DOCX to HWPX
        Returns: (hwpx_bytes, stats_dict)
        """
        if not os.path.exists(pdf_file_path):
            raise FileNotFoundError(f"PDF file not found: {pdf_file_path}")

        # Gather PDF basic stats (pages, quick text preview)
        pdf_doc = fitz.open(pdf_file_path)
        total_pages = len(pdf_doc)
        total_chars = 0
        preview_snippets = []

        for p_idx, page in enumerate(pdf_doc):
            p_text = page.get_text()
            total_chars += len(p_text.strip())
            if p_idx < 3 and p_text.strip():
                preview_snippets.append(f"--- [페이지 {p_idx + 1}] ---\n" + p_text.strip()[:600])
        pdf_doc.close()

        preview_text = "\n\n".join(preview_snippets)

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_docx_path = os.path.join(temp_dir, "intermediate.docx")
            temp_hwpx_path = os.path.join(temp_dir, "result.hwpx")

            # ── Step 1: PDF -> DOCX Layout Reconstruction ────────
            cv = Converter(pdf_file_path)
            cv.convert(temp_docx_path, start=0, end=None)
            cv.close()

            # ── Step 2: DOCX -> HWPX Standard Package Generation ──
            hwpx_bytes = convert_docx_to_hwpx(temp_docx_path, temp_hwpx_path)

            if hwpx_out_path:
                with open(hwpx_out_path, 'wb') as f_out:
                    f_out.write(hwpx_bytes)

            # Extract detailed statistics from the generated intermediate docx
            import docx
            doc = docx.Document(temp_docx_path)
            total_paragraphs = len(doc.paragraphs)
            total_tables = len(doc.tables)

            stats = {
                "totalPages": total_pages,
                "totalParagraphs": total_paragraphs,
                "totalTables": total_tables,
                "totalChars": total_chars,
                "fileSizeBytes": len(hwpx_bytes),
                "previewText": preview_text
            }

            return hwpx_bytes, stats

